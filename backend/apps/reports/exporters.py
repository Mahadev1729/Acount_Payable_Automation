"""
Report Exporters
Generate PDF, Excel, and CSV exports for all report types.
"""
import io
import csv
import datetime
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, HRFlowable
from reportlab.lib.enums import TA_CENTER


def export_to_csv(data, headers):
    """
    Export data to CSV format.
    Returns BytesIO buffer.
    """
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=headers, extrasaction='ignore')
    writer.writeheader()
    writer.writerows(data)
    output = io.BytesIO()
    output.write(buffer.getvalue().encode('utf-8-sig'))  # BOM for Excel compatibility
    output.seek(0)
    return output


def export_to_excel(data, headers, sheet_name='Report'):
    """
    Export data to Excel format using openpyxl.
    Returns BytesIO buffer.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    # Header style
    header_fill = PatternFill(start_color='1e40af', end_color='1e40af', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True, size=10)
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

    # Write headers
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align

    # Write data
    for row_idx, row_data in enumerate(data, 2):
        for col_idx, header in enumerate(headers, 1):
            value = row_data.get(header, '')
            ws.cell(row=row_idx, column=col_idx, value=value)

        # Alternating row colors
        if row_idx % 2 == 0:
            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = PatternFill(
                    start_color='f0f4ff', end_color='f0f4ff', fill_type='solid'
                )

    # Auto-width columns
    for col_idx, header in enumerate(headers, 1):
        max_length = max(len(str(header)), 10)
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_length + 2, 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def export_to_pdf(title, data, headers, landscape_mode=False):
    """
    Export data to PDF using ReportLab.
    Returns BytesIO buffer.
    """
    buffer = io.BytesIO()
    pagesize = landscape(A4) if landscape_mode else A4
    doc = SimpleDocTemplate(buffer, pagesize=pagesize,
                             rightMargin=1.5 * cm, leftMargin=1.5 * cm,
                             topMargin=2 * cm, bottomMargin=2 * cm)

    styles = getSampleStyleSheet()
    story = []

    # Title
    title_style = ParagraphStyle('Title', parent=styles['Title'], fontSize=16,
                                  textColor=colors.HexColor('#1e40af'), spaceAfter=4)
    story.append(Paragraph(f"AP AUTOMATION SYSTEM", title_style))
    story.append(Paragraph(title, ParagraphStyle('Sub', parent=styles['Normal'],
                                                   fontSize=11, textColor=colors.HexColor('#6b7280'))))
    story.append(Paragraph(f"Generated: {datetime.datetime.now().strftime('%d %b %Y, %I:%M %p')}",
                            ParagraphStyle('Date', parent=styles['Normal'], fontSize=8,
                                            textColor=colors.HexColor('#9ca3af'))))
    story.append(HRFlowable(width="100%", thickness=2, color=colors.HexColor('#1e40af')))
    story.append(Spacer(1, 0.2 * cm))

    # Table
    table_data = [headers]
    for row in data:
        table_data.append([str(row.get(h, '')) for h in headers])

    col_width = (pagesize[0] - 3 * cm) / len(headers)
    t = Table(table_data, colWidths=[col_width] * len(headers), repeatRows=1)
    t.setStyle(TableStyle([
        # Header
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e40af')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 8),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('PADDING', (0, 0), (-1, -1), 5),
        # Data rows
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 7.5),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#f8faff'), colors.white]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#d1d5db')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))

    story.append(t)
    doc.build(story)
    buffer.seek(0)
    return buffer
